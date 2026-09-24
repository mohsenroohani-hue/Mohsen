"""ساخت فایل‌های اکسل از رکوردهای استخراج‌شده.

محدودیت اکسل: هر سلول حداکثر ۳۲٬۷۶۷ نویسه می‌پذیرد. متن‌های بلندتر به چند
«بخش» شکسته می‌شوند و در ردیف‌های پیاپی می‌آیند تا هیچ بخشی از متن حذف نشود.
"""

from __future__ import annotations

import os
import re
from collections import Counter, defaultdict
from typing import Dict, List, Optional, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from . import cities as cities_mod
from . import jalali
from .scraper import Record

EXCEL_CELL_LIMIT = 32767
CHUNK_SIZE = 30000          # حاشیه امن نسبت به سقف اکسل
NO_CITY_LABEL = "بدون شهر (سراسری/عمومی)"

_ILLEGAL_XML = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F]")

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(bold=True, size=13, color="1F4E79")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def sanitize(value) -> str:
    """نویسه‌های غیرمجاز XML را حذف می‌کند تا اکسل خطا ندهد."""
    if value is None:
        return ""
    text = str(value)
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    return _ILLEGAL_XML.sub("", text)


def chunk_text(text: str, size: int = CHUNK_SIZE) -> List[str]:
    """متن را به قطعات زیر سقف سلول اکسل می‌شکند (ترجیحاً روی مرز خط)."""
    text = sanitize(text)
    if not text:
        return [""]
    if len(text) <= size:
        return [text]
    chunks: List[str] = []
    start = 0
    while start < len(text):
        end = min(start + size, len(text))
        if end < len(text):
            window = text.rfind("\n", start + int(size * 0.6), end)
            if window == -1:
                window = text.rfind(" ", start + int(size * 0.6), end)
            if window > start:
                end = window + 1
        chunks.append(text[start:end])
        start = end
    return chunks


# ------------------------------------------------------------------ قالب‌بندی
def style_sheet(ws: Worksheet, widths: Sequence[int], freeze: str = "A2",
                wrap_cols: Sequence[int] = ()) -> None:
    ws.sheet_view.rightToLeft = True
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[1].height = 30
    if freeze:
        ws.freeze_panes = freeze
    if ws.max_row >= 1 and ws.max_column >= 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    for col in wrap_cols:
        letter = get_column_letter(col)
        for row in range(2, ws.max_row + 1):
            ws[f"{letter}{row}"].alignment = Alignment(
                wrap_text=True, vertical="top", horizontal="right")


def append_rows(ws: Worksheet, header: Sequence[str],
                rows: Sequence[Sequence]) -> None:
    ws.append(list(header))
    for row in rows:
        ws.append([sanitize(v) if isinstance(v, str) else v for v in row])


# ------------------------------------------------------------------ مرتب‌سازی
def sort_by_date(records: List[Record]) -> List[Record]:
    return sorted(records, key=lambda r: (
        jalali.sort_key(jalali.parse_jalali_date(r.date_jalali)), r.title))


def sort_by_city(records: List[Record]) -> List[Record]:
    """مرتب‌سازی بر پایه الفبای فارسی شهر، سپس تاریخ؛ بی‌شهرها در انتها."""
    return sorted(records, key=lambda r: (
        0 if r.city else 1,
        cities_mod.persian_sort_key(r.city) if r.city else [],
        jalali.sort_key(jalali.parse_jalali_date(r.date_jalali)),
        cities_mod.persian_sort_key(r.title)))


def city_label(rec: Record) -> str:
    return rec.city or NO_CITY_LABEL


# ------------------------------------------------------------------ فهرست
INDEX_HEADER = ["ردیف", "شناسه", "عنوان مصوبه", "شهر", "استان",
                "تاریخ تصویب (شمسی)", "سال", "ماه", "تاریخ میلادی",
                "تعداد نویسه متن", "فایل متن کامل", "پیوند"]
INDEX_WIDTHS = [7, 11, 70, 18, 20, 18, 8, 12, 14, 15, 26, 46]


def _index_row(i: int, rec: Record, text_file: str) -> List:
    return [i, rec.id, rec.title, city_label(rec), rec.province or "",
            rec.date_jalali, rec.year_jalali, rec.month_jalali,
            rec.date_gregorian, rec.body_length, text_file, rec.url]


def build_index_workbook(records: List[Record], out_path: str,
                         file_map: Dict[str, str],
                         text_files: List[str]) -> str:
    wb = Workbook()

    # ---- راهنما
    ws = wb.active
    ws.title = "راهنما"
    ws.sheet_view.rightToLeft = True
    with_city = sum(1 for r in records if r.city)
    dated = sum(1 for r in records if r.date_jalali)
    years = sorted({r.year_jalali for r in records if r.year_jalali})
    info = [
        ("مصوبات شورای عالی شهرسازی و معماری ایران", ""),
        ("", ""),
        ("منبع", "https://rc.majlis.ir/fa/law/search?lu_approve_reference=lar15"),
        ("تعداد کل مصوبات", len(records)),
        ("دارای تاریخ تصویب", dated),
        ("دارای شهر شناسایی‌شده", with_city),
        ("تعداد شهرهای متمایز", len({r.city for r in records if r.city})),
        ("بازه سال‌ها (شمسی)", f"{years[0]} تا {years[-1]}" if years else "—"),
        ("مجموع نویسه‌های متن", sum(r.body_length for r in records)),
        ("", ""),
        ("برگه «بر اساس تاریخ»", "فهرست کامل مرتب‌شده بر حسب تاریخ تصویب"),
        ("برگه «بر اساس شهر»", "همان فهرست، مرتب‌شده بر حسب نام شهر و سپس تاریخ"),
        ("برگه «آمار شهرها»", "تعداد مصوبه هر شهر"),
        ("برگه «آمار سال‌ها»", "تعداد مصوبه در هر سال شمسی"),
        ("", ""),
        ("فایل‌های متن کامل", "، ".join(text_files) if text_files else "—"),
        ("نکته", "متن‌های بلندتر از ۳۰٬۰۰۰ نویسه به چند «بخش» در ردیف‌های پیاپی شکسته شده‌اند."),
        ("نکته", "ستون «فایل متن کامل» در فهرست، فایل حاوی متن هر مصوبه را نشان می‌دهد."),
    ]
    for key, val in info:
        ws.append([sanitize(key), sanitize(val) if isinstance(val, str) else val])
    ws["A1"].font = TITLE_FONT
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 78
    label_font = Font(bold=True)
    for row in range(1, ws.max_row + 1):
        ws[f"A{row}"].font = TITLE_FONT if row == 1 else label_font
        ws[f"B{row}"].alignment = Alignment(wrap_text=True, vertical="top", horizontal="right")

    # ---- بر اساس تاریخ
    ws = wb.create_sheet("بر اساس تاریخ")
    by_date = sort_by_date(records)
    append_rows(ws, INDEX_HEADER,
                [_index_row(i, r, file_map.get(r.id, "")) for i, r in enumerate(by_date, 1)])
    style_sheet(ws, INDEX_WIDTHS, wrap_cols=(3,))

    # ---- بر اساس شهر
    ws = wb.create_sheet("بر اساس شهر")
    by_city = sort_by_city(records)
    append_rows(ws, INDEX_HEADER,
                [_index_row(i, r, file_map.get(r.id, "")) for i, r in enumerate(by_city, 1)])
    style_sheet(ws, INDEX_WIDTHS, wrap_cols=(3,))

    # ---- آمار شهرها
    ws = wb.create_sheet("آمار شهرها")
    grouped: Dict[str, List[Record]] = defaultdict(list)
    for r in records:
        grouped[city_label(r)].append(r)
    rows = []
    for name in sorted(grouped, key=lambda n: (n == NO_CITY_LABEL, -len(grouped[n]),
                                                cities_mod.persian_sort_key(n))):
        group = grouped[name]
        dates = sorted(d for d in (g.date_jalali for g in group) if d)
        rows.append([name, group[0].province or "", len(group),
                     dates[0] if dates else "", dates[-1] if dates else ""])
    append_rows(ws, ["شهر", "استان", "تعداد مصوبه", "نخستین تاریخ", "آخرین تاریخ"], rows)
    style_sheet(ws, [30, 22, 14, 18, 18])

    # ---- آمار سال‌ها
    ws = wb.create_sheet("آمار سال‌ها")
    counter = Counter(r.year_jalali or "نامشخص" for r in records)
    rows = sorted(((y, c) for y, c in counter.items()),
                  key=lambda x: (x[0] == "نامشخص", x[0]))
    append_rows(ws, ["سال شمسی", "تعداد مصوبه"], list(rows))
    style_sheet(ws, [16, 16])

    wb.save(out_path)
    return out_path


# ------------------------------------------------------------------ متن کامل
TEXT_HEADER = ["ردیف", "شناسه", "عنوان مصوبه", "شهر", "استان",
               "تاریخ تصویب (شمسی)", "سال", "تاریخ میلادی",
               "شماره بخش", "تعداد بخش", "متن مصوبه", "پیوند"]
TEXT_WIDTHS = [7, 11, 55, 16, 18, 16, 8, 13, 11, 11, 120, 40]


def build_text_workbooks(records: List[Record], out_dir: str, base_name: str,
                         per_file: int = 150) -> (List[str], Dict[str, str]):
    """متن کامل را در چند فایل اکسل می‌نویسد و نگاشت شناسه→فایل را برمی‌گرداند."""
    os.makedirs(out_dir, exist_ok=True)
    ordered = sort_by_date(records)
    files: List[str] = []
    file_map: Dict[str, str] = {}

    total_files = max(1, (len(ordered) + per_file - 1) // per_file)
    for part in range(total_files):
        batch = ordered[part * per_file:(part + 1) * per_file]
        if not batch and part > 0:
            break
        name = f"{base_name}_متن_کامل_{part + 1:02d}_از_{total_files:02d}.xlsx"
        path = os.path.join(out_dir, name)

        wb = Workbook()
        ws = wb.active
        ws.title = "متن کامل"
        ws.append(TEXT_HEADER)
        row_no = 0
        for rec in batch:
            file_map[rec.id] = name
            chunks = chunk_text(rec.body)
            for pos, chunk in enumerate(chunks, 1):
                row_no += 1
                ws.append([row_no, rec.id, sanitize(rec.title), city_label(rec),
                           rec.province or "", rec.date_jalali, rec.year_jalali,
                           rec.date_gregorian, pos, len(chunks), chunk, rec.url])
        style_sheet(ws, TEXT_WIDTHS, wrap_cols=(3, 11))
        for row in range(2, ws.max_row + 1):
            ws.row_dimensions[row].height = 95
        wb.save(path)
        files.append(name)
    return files, file_map


# ------------------------------------------------------------------ خروجی کامل
def export_all(records: List[Record], out_dir: str,
               base_name: str = "مصوبات_شورای_عالی_شهرسازی",
               per_file: int = 150,
               write_txt: bool = True) -> Dict[str, object]:
    """همه خروجی‌ها را می‌سازد و گزارش می‌دهد."""
    os.makedirs(out_dir, exist_ok=True)
    text_files, file_map = build_text_workbooks(records, out_dir, base_name, per_file)
    index_path = os.path.join(out_dir, f"{base_name}_فهرست.xlsx")
    build_index_workbook(records, index_path, file_map, text_files)

    txt_dir = ""
    if write_txt:
        txt_dir = os.path.join(out_dir, "متن_خام")
        os.makedirs(txt_dir, exist_ok=True)
        for rec in records:
            safe_city = re.sub(r"[^\w؀-ۿ]+", "_", city_label(rec))[:30]
            fname = f"{rec.date_jalali.replace('/', '-') or '0000-00-00'}_{safe_city}_{rec.id}.txt"
            with open(os.path.join(txt_dir, fname), "w", encoding="utf-8") as fh:
                fh.write(f"{rec.title}\n")
                fh.write(f"شهر: {city_label(rec)}   استان: {rec.province}\n")
                fh.write(f"تاریخ: {rec.date_jalali} ({rec.date_gregorian})\n")
                fh.write(f"پیوند: {rec.url}\n")
                fh.write("-" * 70 + "\n")
                fh.write(rec.body)

    return {
        "index_file": os.path.basename(index_path),
        "text_files": text_files,
        "txt_dir": txt_dir,
        "records": len(records),
        "total_chars": sum(r.body_length for r in records),
    }
