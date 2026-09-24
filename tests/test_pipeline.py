"""آزمون سرتاسری خط لوله با HTML ساختگیِ شبیه به سایت مرجع."""

import os
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import load_workbook

from majlis_lar15 import cities, excel_builder, jalali, scraper

SEARCH_HTML = """
<html><head><title>جستجو</title></head><body>
<div class="header"><a href="/fa/home">خانه</a></div>
<table class="results">
  <tr><td><a href="/fa/law/show/802645">مصوبه شورای عالی شهرسازی و معماری ایران در خصوص طرح جامع شهر بندرعباس</a></td>
      <td>۱۳۹۸/۰۵/۱۲</td></tr>
  <tr><td><a href="/fa/law/show/802646">مصوبه شورای عالی شهرسازی و معماری ایران پیرامون طرح تفصیلی شهر قائم‌شهر</a></td>
      <td>۱۳۹۰/۱۱/۰۳</td></tr>
  <tr><td><a href="/fa/law/show/802647">مصوبه شورای عالی شهرسازی و معماری ایران درباره ضوابط عام ساخت‌وساز</a></td>
      <td>۱۴۰۱/۰۲/۲۵</td></tr>
</table>
<div class="pager">
  <a href="?lu_approve_reference=lar15&page=1">1</a>
  <a href="?lu_approve_reference=lar15&page=2">2</a>
  <a href="?lu_approve_reference=lar15&page=3">3</a>
</div>
</body></html>
"""

DETAIL_HTML = """
<html><head><title>مصوبه | مرکز پژوهش‌های مجلس</title></head><body>
<nav><a href="/fa/home">خانه</a><a href="/fa/law">قوانین</a></nav>
<h1>مصوبه شورای عالی شهرسازی و معماری ایران در خصوص طرح جامع شهر بندرعباس</h1>
<div class="meta">شماره: 300/12345 &nbsp; تاریخ تصویب: ۱۳۹۸/۰۵/۱۲ &nbsp;
مرجع تصویب: شورای عالی شهرسازی و معماری ایران</div>
<div id="law-content">
<p>شورای عالی شهرسازی و معماری ایران در جلسه مورخ ۱۳۹۸/۰۵/۱۲ طرح جامع شهر بندرعباس
واقع در استان هرمزگان را مورد بررسی قرار داد و مصوب نمود:</p>
<p>۱- محدوده شهر بندرعباس مطابق نقشه پیوست به مساحت ۱۲۳۴ هکتار تعیین می‌گردد.</p>
<p>۲- حریم شهر بندرعباس مطابق ضوابط ابلاغی رعایت شود.</p>
<p>۳- دبیرخانه شورای عالی موظف به ابلاغ این مصوبه است.</p>
</div>
<footer>کلیه حقوق محفوظ است</footer>
</body></html>
"""

LONG_DETAIL_HTML = DETAIL_HTML.replace(
    "<p>۳- دبیرخانه", "<p>" + ("متن بسیار بلند آزمایشی برای بررسی شکستن سلول. " * 2000) + "</p><p>۳- دبیرخانه")


class TestJalali(unittest.TestCase):
    def test_known_conversions(self):
        self.assertEqual(jalali.format_gregorian((1398, 5, 12)), "2019-08-03")
        self.assertEqual(jalali.format_gregorian((1403, 1, 1)), "2024-03-20")
        self.assertEqual(jalali.format_gregorian((1357, 11, 22)), "1979-02-11")

    def test_parsing_variants(self):
        self.assertEqual(jalali.parse_jalali_date("تاریخ تصویب ۱۳۹۸/۰۵/۱۲"), (1398, 5, 12))
        self.assertEqual(jalali.parse_jalali_date("مصوب ۱۲ مرداد ۱۳۹۸"), (1398, 5, 12))
        self.assertEqual(jalali.parse_jalali_date("مورخ 1385-2-3"), (1385, 2, 3))
        self.assertIsNone(jalali.parse_jalali_date("بدون هیچ تاریخی"))

    def test_leap_years(self):
        self.assertTrue(jalali.valid_jalali(1399, 12, 30))
        self.assertFalse(jalali.valid_jalali(1398, 12, 30))


class TestCities(unittest.TestCase):
    def test_extraction(self):
        self.assertEqual(cities.extract_city("طرح جامع شهر بندرعباس")[0], "بندرعباس")
        self.assertEqual(cities.extract_city("طرح تفصیلی شهر قائم‌شهر")[0], "قائم شهر")
        self.assertEqual(cities.extract_city("ضوابط عام ساخت‌وساز")[0], "")

    def test_persian_collation(self):
        names = ["یزد", "پردیس", "اهواز", "کرج", "گرگان", "چابهار", "تهران"]
        ordered = sorted(names, key=cities.persian_sort_key)
        self.assertEqual(ordered,
                         ["اهواز", "پردیس", "تهران", "چابهار", "کرج", "گرگان", "یزد"])
        # ترتیب خام یونیکد نادرست است و نباید با ترتیب الفبایی یکی باشد
        self.assertNotEqual(sorted(names), ordered)

    def test_province(self):
        self.assertEqual(cities.extract_province("طرح جامع شهر بندرعباس", city="بندرعباس"),
                         "هرمزگان")
        self.assertEqual(cities.extract_province("مصوبه در استان گیلان"), "گیلان")


class TestParsing(unittest.TestCase):
    def test_search_page(self):
        rows = scraper.parse_search_page(SEARCH_HTML)
        self.assertEqual(len(rows), 3)
        self.assertEqual(rows[0]["id"], "802645")
        self.assertIn("بندرعباس", rows[0]["title"])
        self.assertEqual(scraper.max_page_number(SEARCH_HTML), 3)

    def test_detail_page(self):
        rec = scraper.parse_detail_page(DETAIL_HTML, "802645")
        self.assertIn("بندرعباس", rec.title)
        self.assertNotIn("مرکز پژوهش", rec.title)
        self.assertEqual(rec.date_jalali, "1398/05/12")
        self.assertEqual(rec.date_gregorian, "2019-08-03")
        self.assertEqual(rec.city, "بندرعباس")
        self.assertEqual(rec.province, "هرمزگان")
        self.assertEqual(rec.document_number, "300/12345")
        self.assertIn("۱۲۳۴ هکتار", rec.body)
        self.assertNotIn("کلیه حقوق محفوظ", rec.body)   # فوتر حذف شده
        self.assertNotIn("خانه", rec.body.split("\n")[0])
        self.assertEqual(rec.parse_warnings, [])


class TestChunking(unittest.TestCase):
    def test_lossless(self):
        text = "الف ب پ " * 20000
        chunks = excel_builder.chunk_text(text)
        self.assertGreater(len(chunks), 1)
        self.assertEqual("".join(chunks), text)
        self.assertTrue(all(len(c) <= excel_builder.EXCEL_CELL_LIMIT for c in chunks))

    def test_sanitize_removes_control_chars(self):
        self.assertEqual(excel_builder.sanitize("a\x00b\x07c"), "abc")


class TestEndToEnd(unittest.TestCase):
    def test_full_export(self):
        recs = [
            scraper.parse_detail_page(DETAIL_HTML, "802645"),
            scraper.parse_detail_page(
                DETAIL_HTML.replace("بندرعباس", "قائم‌شهر")
                           .replace("۱۳۹۸/۰۵/۱۲", "۱۳۹۰/۱۱/۰۳")
                           .replace("هرمزگان", "مازندران"), "802646"),
            scraper.parse_detail_page(LONG_DETAIL_HTML, "802647"),
            scraper.parse_detail_page(
                DETAIL_HTML.replace("بندرعباس", "پردیس").replace("هرمزگان", "تهران"),
                "802648"),
            scraper.parse_detail_page(
                DETAIL_HTML.replace(
                    "در خصوص طرح جامع شهر بندرعباس", "پیرامون ضوابط عام ساخت‌وساز")
                           .replace("بندرعباس", "کشور").replace("هرمزگان", "کشور"),
                "802649"),
        ]
        with tempfile.TemporaryDirectory() as tmp:
            report = excel_builder.export_all(recs, tmp, per_file=2)
            self.assertEqual(report["records"], 5)
            self.assertEqual(len(report["text_files"]), 3)   # تقسیم به سه فایل

            idx = load_workbook(os.path.join(tmp, str(report["index_file"])))
            self.assertIn("بر اساس تاریخ", idx.sheetnames)
            self.assertIn("بر اساس شهر", idx.sheetnames)
            wsc = idx["بر اساس شهر"]
            city_col = [wsc.cell(r, 4).value for r in range(2, wsc.max_row + 1)]
            named = [c for c in city_col if c != excel_builder.NO_CITY_LABEL]
            self.assertEqual(named, sorted(named, key=cities.persian_sort_key))
            self.assertEqual(city_col[-1], excel_builder.NO_CITY_LABEL)

            ws = idx["بر اساس تاریخ"]
            dates = [ws.cell(r, 6).value for r in range(2, ws.max_row + 1)]
            self.assertEqual(dates, sorted(dates))          # ترتیب زمانی درست
            self.assertEqual(ws.max_row, 6)                 # سرستون + ۵ رکورد

            # متن کامل: بازسازی بدون افت
            rebuilt = {}
            for name in report["text_files"]:
                wsx = load_workbook(os.path.join(tmp, name))["متن کامل"]
                for r in range(2, wsx.max_row + 1):
                    law_id = str(wsx.cell(r, 2).value)
                    rebuilt.setdefault(law_id, "")
                    rebuilt[law_id] += wsx.cell(r, 11).value or ""
            for rec in recs:
                self.assertEqual(rebuilt[rec.id], excel_builder.sanitize(rec.body),
                                 f"متن مصوبه {rec.id} کامل بازسازی نشد")

            # فایل‌های متنی پشتیبان
            txts = os.listdir(os.path.join(tmp, "متن_خام"))
            self.assertEqual(len(txts), 5)


if __name__ == "__main__":
    unittest.main(verbosity=2)
