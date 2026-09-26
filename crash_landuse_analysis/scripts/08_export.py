"""Export GIS layers (GeoPackage for QGIS / ArcGIS Pro / GeoDa), GeoDa weights files,
a data dictionary and an Excel workbook with every result table."""
import numpy as np
import pandas as pd
import geopandas as gpd
from libpysal.weights import KNN
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

import common as C

GPKG = C.GIS / "crash_landuse_spatial.gpkg"
if GPKG.exists():
    GPKG.unlink()

typ = pd.read_pickle(C.OUT / "typology_sites.pkl")
mg_path = C.OUT / "mgwr_sites.pkl"
mg = pd.read_pickle(mg_path) if mg_path.exists() else {}
dictionary = []

KEEP = ["SEGMENTID", "COUNTY", "region", "ROUTENAME", "Route_grp", "SEGTYPE", "SEGLENGTH",
        "RN_AADT_AVG", "RN_LANES_AVG_TOT", "RN_MAXSPEED_AVG", "raised_median", "RN_SHLDO_AV_TWIDTH",
        "bikelane", "Busstops", "CN_POPDEN_SQMI", "CN_MED_HH_INC", "exposure"]
for fac in ("int", "seg"):
    g = C.load(fac)
    esda_ = pd.read_pickle(C.OUT / f"esda_sites_{fac}.pkl")
    cols = [c for c in KEEP if c in g] + (["AADTCrossSt"] if fac == "int" else [])
    cols += list(C.LAND_USES)
    ycols = [f"y_{t}_{s}" for t in C.CRASH_TYPES for s in C.SEVERITIES]
    df = g[cols + ycols].copy()
    df = df.merge(esda_, on="SEGMENTID", how="left").merge(typ[fac], on="SEGMENTID", how="left")
    if fac in mg:
        df = df.merge(mg[fac], on="SEGMENTID", how="left")
    for c in df.columns:
        if df[c].dtype == object:
            df[c] = df[c].astype(str)
    name = "intersections" if fac == "int" else "segments"
    # line geometry (intersection influence area / segment)
    gpd.GeoDataFrame(df, geometry=g.geometry.values, crs=C.CRS).to_file(GPKG, layer=f"{name}_lines", driver="GPKG")
    # point geometry (mid-point) - use this layer in GeoDa and for point-pattern tools
    gpd.GeoDataFrame(df, geometry=g["pt"].values, crs=C.CRS).to_file(GPKG, layer=f"{name}_points", driver="GPKG")
    df.assign(wkt=g["pt"].to_wkt().values).to_csv(C.GIS / f"{name}_points_with_results.csv", index=False)

    # GeoDa-format GAL weights (k = 6 nearest neighbours, ids = SEGMENTID)
    w = KNN.from_array(np.column_stack([g["pt"].x, g["pt"].y]), k=6)
    ids = g["SEGMENTID"].tolist()
    with open(C.GIS / f"{name}_knn6.gal", "w") as fh:
        fh.write(f"0 {len(ids)} {name}_points SEGMENTID\n")
        for i, nb in w.neighbors.items():
            fh.write(f"{ids[i]} {len(nb)}\n")
            fh.write(" ".join(str(ids[j]) for j in nb) + "\n")

    for c in df.columns:
        if c in dictionary and False:
            continue
    if fac == "int":
        for c in df.columns:
            desc = ""
            if c.startswith("y_"):
                _, t, s = c.split("_", 2)
                desc = f"{C.CRASH_TYPES[t]} crashes, {C.SEV_LABEL[s]} (4-year count)"
            elif c.startswith("eb_"):
                desc = "Empirical-Bayes smoothed crash rate (per MEV at intersections, per 100 MVMT on segments): " + c[3:]
            elif c.startswith("lisa_"):
                desc = "LISA cluster class (EB rate, k=6, 999 perms, p<0.05): " + c[5:]
            elif c.startswith("lisap_"):
                desc = "LISA pseudo p-value: " + c[6:]
            elif c.startswith("giz_"):
                desc = "Getis-Ord Gi* z-score (EB rate): " + c[4:]
            elif c.startswith("gi_"):
                desc = "Gi* hot/cold spot class (90/95/99%): " + c[3:]
            elif c.startswith("ljccls_"):
                desc = "Local join count class (binary: any crash of this type/severity): " + c[7:]
            elif c.startswith("ljcp_"):
                desc = "Local join count pseudo p-value: " + c[5:]
            elif c.startswith("ljc_"):
                desc = "Local join count statistic: " + c[4:]
            elif c.startswith("bvlisa_"):
                desc = "Bivariate LISA class: land use at site x spatial lag of EB crash rate: " + c[7:]
            elif c.startswith("mgwr_b_"):
                desc = "MGWR local standardised coefficient (outcome ln EB KAB rate): " + c[7:]
            elif c.startswith("mgwr_sig_"):
                desc = "MGWR local significance (corrected critical t): " + c[9:]
            elif c in C.LAND_USES:
                desc = C.LAND_USES[c] + " (count)"
            elif c == "typology":
                desc = "Land-use context typology (K-means, k=5)"
            elif c == "exposure":
                desc = "Exposure: million entering vehicles (intersections) / 100 million VMT (segments), 4 years"
            dictionary.append({"field": c, "description": desc})

pd.DataFrame(dictionary).to_csv(C.GIS / "data_dictionary.csv", index=False)
cty = gpd.read_file(C.DATA / "fl_counties.gpkg")
cty.to_file(GPKG, layer="florida_counties", driver="GPKG")
print("gpkg written:", GPKG)

# ------------------------------------------------------------------ Excel workbook
sheets = [
    ("README", None),
    ("crash_counts", "t_crash_counts.csv"),
    ("descriptives", "t_descriptives.csv"),
    ("landuse_prevalence", "t_landuse_prevalence.csv"),
    ("global_moran", "t_global_moran.csv"),
    ("incremental_moran", "t_incremental_moran.csv"),
    ("bivariate_moran_each_LU", "t_bivariate_moran.csv"),
    ("hotspot_share_corridor", "t_hotspot_share_by_corridor.csv"),
    ("typology_k_selection", "t_typology_kselection.csv"),
    ("typology_profiles", "t_typology_profiles.csv"),
    ("typology_rates", "t_typology_rates.csv"),
    ("typology_IRR", "res_typology_irr.csv"),
    ("IRR_each_LU_single", "res_single_landuse.csv"),
    ("IRR_all_LU_joint", "res_joint_landuse.csv"),
    ("joint_model_controls", "res_joint_controls.csv"),
    ("model_diagnostics", "res_model_diagnostics.csv"),
    ("LU_exacerbation_summary", "t_landuse_exacerbation_summary.csv"),
    ("MGWR_coefficients", "res_mgwr_coefficients.csv"),
    ("MGWR_fit", "res_mgwr_fit.csv"),
]
readme = [
    ["Land use x crash type x severity: spatial analysis results"],
    [""],
    ["Facilities", "int = intersections (n=489), seg = segments (n=334); six Florida counties; 4-year crash totals"],
    ["Severity", "TOT = all crashes; KAB = fatal + incapacitating + non-incapacitating injury; KSI = fatal + incapacitating; FAT = fatal"],
    ["Crash types", ", ".join(f"{k} = {v}" for k, v in C.CRASH_TYPES.items())],
    ["Land uses", ", ".join(f"{k} = {v}" for k, v in C.LAND_USES.items())],
    ["IRR", "Incidence rate ratio per additional establishment (per 10 for mixed-use and office buildings); counts winsorised at the 99th percentile"],
    ["Estimators", "NB2 = negative binomial with full controls + county fixed effects (outcomes with >=150 crashes, non-fatal); Firth-Poisson = bias-reduced Poisson with quasi-Poisson SEs and reduced controls + region fixed effects (sparse outcomes, all fatal models)"],
    ["Spatial filter", "Moran eigenvectors (k=6 KNN) selected on residuals and added to every model for that outcome (n_eigvec)"],
    ["q", "Benjamini-Hochberg false-discovery-rate adjusted p-value within facility x severity"],
    ["single vs joint", "single = one land use at a time (plus all controls); joint = all land uses entered together"],
]
wb = Workbook()
hdr_fill = PatternFill("solid", fgColor="DCE8F7")
for name, fn in sheets:
    ws = wb.active if name == "README" else wb.create_sheet(name)
    ws.title = name
    if fn is None:
        for r in readme:
            ws.append(r)
        ws["A1"].font = Font(bold=True, size=13)
        ws.column_dimensions["A"].width = 18
        ws.column_dimensions["B"].width = 140
        for row in ws.iter_rows(min_row=3):
            row[0].font = Font(bold=True)
            if len(row) > 1:
                row[1].alignment = Alignment(wrap_text=True, vertical="top")
        continue
    p = C.TAB / fn
    if not p.exists():
        ws.append([f"{fn} not available"])
        continue
    df = pd.read_csv(p)
    ws.append(list(df.columns))
    for row in df.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and np.isnan(v)) else
                   (round(v, 5) if isinstance(v, float) else v) for v in row])
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = hdr_fill
    ws.freeze_panes = "A2"
    for i, col in enumerate(df.columns, 1):
        width = max(len(str(col)), *(len(str(v)) for v in df[col].head(200))) if len(df) else len(col)
        ws.column_dimensions[get_column_letter(i)].width = min(max(9, width + 2), 45)
    ws.auto_filter.ref = ws.dimensions
wb.save(C.TAB / "crash_landuse_results.xlsx")
print("excel written")
